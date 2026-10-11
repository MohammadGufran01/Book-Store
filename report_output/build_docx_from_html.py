"""
Converts the generated report pages into a beautifully formatted Microsoft Word (.docx) document
using python-docx. Handles typography, borders, tables, code boxes, images, and captions.
"""

import os
import re
import html
from html.parser import HTMLParser
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(BASE_DIR, '..'))
DOCX_PATH = os.path.join(BASE_DIR, "Readers_Haven_Project_Report.docx")
HTML_PATH = os.path.join(BASE_DIR, "Readers_Haven_Project_Report.html")

# Image lookup map from exact alt text
ALT_MAP = {
    'Rizvi College Logo': os.path.join(PROJECT_DIR, 'report_assets', 'rId8.jpg'),
    'Plagiarism Summary Graphic': os.path.join(BASE_DIR, 'diagrams', 'fig_plagiarism.png'),
    'Spiral Model Diagram': os.path.join(BASE_DIR, 'diagrams', 'spiral_model.png'),
    'System Architecture Diagram': os.path.join(BASE_DIR, 'diagrams', 'architecture_diagram.png'),
    'Level 0 Data Flow Diagram': os.path.join(BASE_DIR, 'diagrams', 'dfd_level_0.png'),
    'Level 1 Data Flow Diagram': os.path.join(BASE_DIR, 'diagrams', 'dfd_level_1.png'),
    'Activity Diagram': os.path.join(BASE_DIR, 'diagrams', 'activity_diagram.png'),
    'Entity Relationship Diagram': os.path.join(BASE_DIR, 'diagrams', 'er_diagram.png'),
    'Use-Case Diagram': os.path.join(BASE_DIR, 'diagrams', 'use_case_diagram.png'),
    'System Flowchart': os.path.join(BASE_DIR, 'diagrams', 'flowchart_pipeline.png'),
    'Login Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_11_login.png'),
    'Registration Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_12_register.png'),
    'Homepage': os.path.join(BASE_DIR, 'screenshots', 'fig7_1_homepage.png'),
    'Catalog Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_2_catalog.png'),
    'Book Detail Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_3_book_detail.png'),
    'For You Recommendations Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_4_recommendations.png'),
    'Shopping Cart Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_5_cart.png'),
    'Checkout Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_6_checkout.png'),
    'Order History Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_7_order_history.png'),
    'Profile Page': os.path.join(BASE_DIR, 'screenshots', 'fig7_8_user_profile.png'),
    'Analytics Dashboard': os.path.join(BASE_DIR, 'screenshots', 'fig7_9_analytics_dashboard.png'),
    'Django Admin Panel': os.path.join(BASE_DIR, 'screenshots', 'fig7_10_django_admin.png'),
}

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, color="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def sanitize_xml(s):
    if not s:
        return ""
    # Strip illegal XML characters (control chars except tab \x09, newline \x0a, carriage return \x0d)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', str(s))

class ReportHTMLParser(HTMLParser):
    def __init__(self, doc):
        super().__init__()
        self.doc = doc
        self.current_tag = None
        self.tag_stack = []
        self.current_p = None
        self.in_code_box = False
        self.in_table = False
        self.table_data = []
        self.current_row = []
        self.current_cell_text = ""
        self.code_text = ""
        self.bold_active = False
        self.italic_active = False
        self.current_cell_text = ""
        self.code_text = ""
        self.bold_active = False
        self.italic_active = False
        
    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        self.tag_stack.append((tag, attrs_dict))
        self.current_tag = tag
        
        cls = attrs_dict.get('class', '')
        
        if tag == 'div' and 'code-box' in cls:
            self.in_code_box = True
            self.code_text = ""
        elif tag == 'pre':
            self.in_code_box = True
            self.code_text = ""
        elif tag == 'table':
            self.in_table = True
            self.table_data = []
        elif tag == 'tr':
            self.current_row = []
        elif tag in ['td', 'th']:
            self.current_cell_text = ""
        elif tag == 'b' or tag == 'strong':
            self.bold_active = True
        elif tag == 'i' or tag == 'em':
            self.italic_active = True
        elif tag == 'img':
            src = attrs_dict.get('src', '')
            alt = attrs_dict.get('alt', '')
            self.add_image_from_src(src, alt)
        elif tag in ['h1', 'h2', 'h3', 'p', 'li']:
            self.start_paragraph(tag, cls)

    def handle_endtag(self, tag):
        if self.tag_stack and self.tag_stack[-1][0] == tag:
            self.tag_stack.pop()
            
        if tag in ['b', 'strong']:
            self.bold_active = False
        elif tag in ['i', 'em']:
            self.italic_active = False
        elif tag in ['div', 'pre'] and self.in_code_box:
            self.in_code_box = False
            if self.code_text.strip():
                self.add_code_table(self.code_text.strip())
            self.code_text = ""
        elif tag in ['td', 'th']:
            self.current_row.append((tag, self.current_cell_text.strip()))
            self.current_cell_text = ""
        elif tag == 'tr':
            if self.current_row:
                self.table_data.append(self.current_row)
                self.current_row = []
        elif tag == 'table':
            self.in_table = False
            if self.table_data:
                self.add_rendered_table(self.table_data)
                self.table_data = []
        elif tag in ['h1', 'h2', 'h3', 'p', 'li']:
            self.current_p = None

    def handle_data(self, data):
        cleaned = sanitize_xml(data)
        if self.in_code_box:
            self.code_text += cleaned
        elif self.in_table:
            self.current_cell_text += cleaned
        elif self.current_p is not None:
            run = self.current_p.add_run(cleaned)
            run.font.name = "Times New Roman"
            if self.bold_active:
                run.font.bold = True
            if self.italic_active:
                run.font.italic = True
        else:
            # text outside known tags, ignore whitespace
            pass

    def start_paragraph(self, tag, cls):
        p = self.doc.add_paragraph()
        self.current_p = p
        if tag == 'h1' or 'chapter-title' in cls:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            self.bold_active = True
        elif tag == 'h2' or 'section-title' in cls:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            self.bold_active = True
        elif tag == 'h3' or 'subsection-title' in cls:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            self.bold_active = True
        elif tag == 'li':
            p.style = 'List Bullet'
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(15)
        elif 'figure-caption' in cls:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(10)
            self.italic_active = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = Pt(15)

    def add_code_table(self, text):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8F9FA")
        set_cell_borders(cell, "CCCCCC", "4")
        set_cell_margins(cell, 80, 80, 120, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(10)
        run = p.add_run(sanitize_xml(text))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_rendered_table(self, data):
        if not data:
            return
        num_rows = len(data)
        num_cols = max(len(row) for row in data)
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        col_width = Inches(6.5 / max(num_cols, 1))
        
        for r_idx, row in enumerate(data):
            for c_idx, (cell_tag, cell_text) in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.width = col_width
                set_cell_borders(cell, "333333" if r_idx == 0 else "CCCCCC", "4")
                set_cell_margins(cell, 60, 60, 100, 100)
                if r_idx == 0 or cell_tag == 'th':
                    set_cell_background(cell, "EAEAEA")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(12)
                run = p.add_run(sanitize_xml(cell_text))
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                if r_idx == 0 or cell_tag == 'th':
                    run.font.bold = True
                    
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(6)

    def add_image_from_src(self, src, alt):
        clean_alt = alt.strip()
        matched_path = ALT_MAP.get(clean_alt)
        if not matched_path:
            for k, v in ALT_MAP.items():
                if k.lower() in clean_alt.lower() or clean_alt.lower() in k.lower():
                    matched_path = v
                    break
        if matched_path and os.path.exists(matched_path):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run()
            width = Inches(1.5) if 'rId8' in matched_path else Inches(5.8)
            r.add_picture(matched_path, width=width)
            print(f"  [DOCX] Embedded image: {clean_alt} -> {os.path.basename(matched_path)}")
        else:
            print(f"  [DOCX WARNING] Image not found for alt: '{clean_alt}'")

def convert_html_pages_to_docx():
    print("Reading generated HTML file...")
    with open(HTML_PATH, 'r', encoding='utf-8', errors='ignore') as f:
        full_html = f.read()

    # Extract all <div class="page-box">...</div>
    page_regex = re.compile(r'<div class="page-box">(.*?)</div>\s*(?=<div class="page-box">|$)', re.DOTALL)
    pages = page_regex.findall(full_html)
    print(f"Found {len(pages)} pages in HTML report.")

    doc = docx.Document()
    # A4 margins (0.75 in)
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)

    for idx, page_html in enumerate(pages):
        print(f"Processing Page {idx + 1}/{len(pages)}...")
        
        # Remove the page-number footer from parsing (we will let Word or doc format handle page flow)
        page_html_clean = re.sub(r'<div class="page-number">.*?</div>', '', page_html)
        
        # Parse page contents
        parser = ReportHTMLParser(doc)
        parser.feed(page_html_clean)
        
        # Add a page break between pages except for the very last page
        if idx < len(pages) - 1:
            doc.add_page_break()

    print(f"Saving Word Document to {DOCX_PATH}...")
    doc.save(DOCX_PATH)
    file_size = os.path.getsize(DOCX_PATH)
    print(f"DOCX created successfully! Size: {file_size} bytes ({file_size/1024/1024:.2f} MB)")

if __name__ == '__main__':
    convert_html_pages_to_docx()
